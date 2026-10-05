"""CSV and Excel import: file → table → dataset (README §2c, step 1)."""

from __future__ import annotations

import csv
import io
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import numpy as np

from propbench.core import Dataset, Phase, Provenance, Quantity, difference_to_si, identify_fluid, to_si
from propbench.importers.mapping import ImportMapping, Role, UncertaintyKind

CSV_SUFFIXES = {".csv", ".txt", ".tsv", ".dat"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}


class ImportFileError(ValueError):
    """The file cannot be read with the given mapping. Messages name the file row and column."""


@dataclass(frozen=True)
class InMemoryFile:
    """File content sent by the UI (no filesystem path needed); ``name`` gives the type and the dataset name."""

    name: str
    data: bytes = field(repr=False)


Source = str | Path | InMemoryFile


def source_suffix(source: Source) -> str:
    return Path(source.name if isinstance(source, InMemoryFile) else source).suffix.lower()


def source_stem(source: Source) -> str:
    return Path(source.name if isinstance(source, InMemoryFile) else source).stem


def source_text(source: Source) -> str:
    if isinstance(source, InMemoryFile):
        return source.data.decode("utf-8-sig")
    return Path(source).read_text(encoding="utf-8-sig")


def source_binary(source: Source) -> Path | io.BytesIO:
    return io.BytesIO(source.data) if isinstance(source, InMemoryFile) else Path(source)


@dataclass(frozen=True)
class Table:
    """Cells of a sheet below a header row. ``row_numbers`` are 1-based file rows, for error messages."""

    headers: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]
    row_numbers: tuple[int, ...]

    def column_index(self, name: str) -> int:
        try:
            return self.headers.index(name)
        except ValueError:
            raise ImportFileError(f"no column '{name}' (columns: {', '.join(repr(h) for h in self.headers)})") from None


def read_csv(path: Source, mapping: ImportMapping) -> Table:
    """Read a delimited text file (delimiter sniffed unless the mapping sets one; UTF-8 with or without BOM)."""
    text = source_text(path)
    delimiter = mapping.delimiter
    if delimiter is None:
        sample = "\n".join(text.splitlines()[:20])
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    lines = list(csv.reader(text.splitlines(), delimiter=delimiter))
    return _table(lines, mapping)


def read_xlsx(path: Source, mapping: ImportMapping) -> Table:
    """Read one sheet of an Excel workbook (the first sheet unless the mapping names one). Cached values are used."""
    from openpyxl import load_workbook  # imported lazily: only needed for Excel files

    workbook = load_workbook(source_binary(path), read_only=True, data_only=True)
    try:
        if mapping.sheet is None:
            sheet = workbook.worksheets[0]
        elif mapping.sheet in workbook.sheetnames:
            sheet = workbook[mapping.sheet]
        else:
            raise ImportFileError(f"no sheet '{mapping.sheet}' (sheets: {', '.join(workbook.sheetnames)})")
        lines = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _table(lines, mapping)


@dataclass(frozen=True)
class Preview:
    """The first rows of a file as text, to set up a mapping. ``sheets`` lists an Excel workbook's sheets."""

    rows: tuple[tuple[str, ...], ...]
    sheets: tuple[str, ...]
    sheet: str | None
    delimiter: str | None


def preview_file(path: Source, sheet: str | None = None, max_rows: int = 30) -> Preview:
    """Raw cells (as text) of the first ``max_rows`` rows of a CSV or Excel file, without any mapping."""
    suffix = source_suffix(path)
    if suffix in CSV_SUFFIXES:
        text = source_text(path)
        sample = "\n".join(text.splitlines()[:20])
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
        lines = list(csv.reader(text.splitlines()[:max_rows], delimiter=delimiter))
        return Preview(tuple(tuple(c.strip() for c in line) for line in lines), (), None, delimiter)
    if suffix in EXCEL_SUFFIXES:
        from openpyxl import load_workbook

        workbook = load_workbook(source_binary(path), read_only=True, data_only=True)
        try:
            names = tuple(workbook.sheetnames)
            if sheet is not None and sheet not in names:
                raise ImportFileError(f"no sheet '{sheet}' (sheets: {', '.join(names)})")
            ws = workbook[sheet] if sheet is not None else workbook.worksheets[0]
            rows = []
            for row in ws.iter_rows(values_only=True, max_row=max_rows):
                rows.append(tuple("" if c is None else str(c) for c in row))
        finally:
            workbook.close()
        return Preview(tuple(rows), names, ws.title, None)
    raise ImportFileError(f"unsupported file type '{suffix}' (use {sorted(CSV_SUFFIXES | EXCEL_SUFFIXES)})")


def read_table(path: Source, mapping: ImportMapping) -> Table:
    suffix = source_suffix(path)
    if suffix in CSV_SUFFIXES:
        return read_csv(path, mapping)
    if suffix in EXCEL_SUFFIXES:
        return read_xlsx(path, mapping)
    raise ImportFileError(f"unsupported file type '{suffix}' (use {sorted(CSV_SUFFIXES | EXCEL_SUFFIXES)})")


def _table(lines: Sequence[Sequence[Any]], mapping: ImportMapping) -> Table:
    if len(lines) < mapping.header_row:
        raise ImportFileError(f"the file has no row {mapping.header_row} for the header")
    headers = tuple("" if h is None else str(h).strip() for h in lines[mapping.header_row - 1])
    rows, numbers = [], []
    for number in range(mapping.data_start_row, len(lines) + 1):
        row = tuple(lines[number - 1])
        if all(_blank(cell) for cell in row):
            continue
        rows.append(row + (None,) * (len(headers) - len(row)))
        numbers.append(number)
    return Table(headers=headers, rows=tuple(rows), row_numbers=tuple(numbers))


def table_to_dataset(
    table: Table,
    mapping: ImportMapping,
    *,
    name: str,
    fluid: str,
    provenance: Provenance | None = None,
) -> Dataset:
    """Convert a table to an SI dataset using the mapping. The fluid is resolved by name, alias or CAS number."""
    if not table.rows:
        raise ImportFileError("the table has no data rows")

    def numbers(role: Role) -> np.ndarray | None:
        spec = mapping.column(role)
        if spec is None:
            return None
        index = table.column_index(spec.column)
        return np.array(
            [_number(row[index], spec.column, n, mapping.decimal_comma) for row, n in _rows(table)], dtype=float
        )

    t_spec, v_spec = mapping.column(Role.TEMPERATURE), mapping.column(Role.VALUE)
    temperature, raw_values = numbers(Role.TEMPERATURE), numbers(Role.VALUE)
    complete = t_spec and v_spec and t_spec.unit and v_spec.unit and temperature is not None and raw_values is not None
    if not complete:  # ImportMapping validation already guarantees this; checked again for the type checker
        raise ImportFileError("the mapping needs temperature and value columns with units")
    values = to_si(raw_values, v_spec.unit, mapping.quantity)

    def state(role: Role, quantity: Quantity) -> np.ndarray | None:
        raw, spec = numbers(role), mapping.column(role)
        return None if raw is None or spec is None or spec.unit is None else to_si(raw, spec.unit, quantity)

    uncertainty = None
    u_spec = mapping.column(Role.UNCERTAINTY)
    raw_u = numbers(Role.UNCERTAINTY)
    if u_spec is not None and raw_u is not None:
        if u_spec.uncertainty_kind is UncertaintyKind.RELATIVE_PERCENT:
            uncertainty = np.abs(values) * raw_u / 100.0
        elif u_spec.uncertainty_kind is UncertaintyKind.RELATIVE_FRACTION:
            uncertainty = np.abs(values) * raw_u
        else:
            uncertainty = difference_to_si(raw_u, u_spec.unit or v_spec.unit, mapping.quantity)

    phase = None
    p_spec = mapping.column(Role.PHASE)
    if p_spec is not None:
        index = table.column_index(p_spec.column)
        phase = tuple(_phase(row[index], p_spec.column, n) for row, n in _rows(table))

    try:
        return Dataset(
            name=name,
            fluid=identify_fluid(fluid),
            quantity=mapping.quantity,
            temperature=to_si(temperature, t_spec.unit, Quantity.TEMPERATURE),
            values=values,
            pressure=state(Role.PRESSURE, Quantity.PRESSURE),
            molar_density=state(Role.MOLAR_DENSITY, Quantity.MOLAR_DENSITY),
            expanded_uncertainty=uncertainty,
            coverage_factor=mapping.coverage_factor,
            phase=phase,
            provenance=provenance or Provenance(),
        )
    except ValueError as exc:
        raise ImportFileError(str(exc)) from exc


def import_file(
    path: Source,
    mapping: ImportMapping,
    *,
    fluid: str,
    name: str | None = None,
    provenance: Provenance | None = None,
) -> Dataset:
    """Read a CSV or Excel file into a dataset. The dataset name defaults to the file name without extension."""
    table = read_table(path, mapping)
    return table_to_dataset(table, mapping, name=name or source_stem(path), fluid=fluid, provenance=provenance)


def _rows(table: Table):
    return zip(table.rows, table.row_numbers, strict=True)


def _blank(cell: Any) -> bool:
    return cell is None or (isinstance(cell, str) and not cell.strip())


def _number(cell: Any, column: str, row: int, decimal_comma: bool) -> float:
    if _blank(cell):
        raise ImportFileError(f"row {row}, column '{column}': empty cell")
    if isinstance(cell, bool):
        raise ImportFileError(f"row {row}, column '{column}': {cell!r} is not a number")
    if isinstance(cell, int | float):
        return float(cell)
    text = str(cell).strip().replace("\u2212", "-")  # Unicode minus sign (U+2212)
    if decimal_comma:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        raise ImportFileError(f"row {row}, column '{column}': '{cell}' is not a number") from None


_PHASE_WORDS = {
    Phase.LIQUID: {"l", "liq", "liquid", "compressed liquid"},
    Phase.VAPOR: {"v", "g", "vap", "vapor", "vapour", "gas"},
    Phase.SUPERCRITICAL: {"sc", "scf", "supercritical", "fluid"},
    Phase.TWO_PHASE: {"2", "two-phase", "two_phase", "saturated"},
}


def _phase(cell: Any, column: str, row: int) -> Phase:
    text = "" if _blank(cell) else str(cell).strip().casefold()
    if not text:
        return Phase.UNKNOWN
    for phase, words in _PHASE_WORDS.items():
        if text in words:
            return phase
    raise ImportFileError(f"row {row}, column '{column}': unknown phase '{cell}'")
