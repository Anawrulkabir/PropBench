"""Reports (README §2, M4): one document from the project's results, as Markdown, Word (python-docx) or PDF (Typst),
plus a reproducibility bundle.

The report spec comes from the app (or a script) and holds results already computed: datasets, consistency,
fitted models with their validation, the comparison table, the locked selection, references and figures (figure
specs, rendered here). The reference models' check values are recomputed while writing, so the report states
that they are reproduced by this installation. Templates set fonts and layout: ``plain``, ``elsevier``, ``acs``.
"""

from __future__ import annotations

import base64
import io
import json
import platform
import tempfile
import zipfile
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

TEMPLATES = {
    "plain": {"font": "Libertinus Serif", "size": 11, "docx_font": "Calibri"},
    "elsevier": {"font": "Libertinus Serif", "size": 10, "docx_font": "Times New Roman"},
    "acs": {"font": "Libertinus Serif", "size": 10, "docx_font": "Arial"},
}
FORMATS = {
    "md": "text/markdown",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "zip": "application/zip",
}


class ReportError(ValueError):
    """The report spec is invalid or a document could not be written."""


def _fmt(v: Any, digits: int = 4) -> str:
    if v is None:
        return "–"
    if isinstance(v, int | float):
        return f"{v:.{digits}g}"
    return str(v)


# --- content: a list of blocks shared by every writer ---


def _versions() -> dict[str, str]:
    import CoolProp
    import numpy
    import scipy

    import propbench

    return {
        "PropBench": propbench.__version__,
        "Python": platform.python_version(),
        "CoolProp": CoolProp.__version__,
        "NumPy": numpy.__version__,
        "SciPy": scipy.__version__,
    }


def check_value_rows() -> list[list[str]]:
    """Every bundled reference model's check values, recomputed now."""
    from propbench.models.reference import load_registry

    rows = []
    for entry in load_registry():
        model = entry.model()
        for c in entry.check_values:
            got = float(model.predict([c.temperature], [c.molar_density])[0])
            dev = abs(got / c.expected - 1)
            rows.append(
                [
                    entry.label,
                    f"{c.temperature:g} K, {c.molar_density / 1000:g} mol/L",
                    f"{c.expected:.7g}",
                    f"{got:.7g}",
                    "yes" if dev <= c.rel_tol else "NO",
                ]
            )
    return rows


def blocks(spec: Mapping[str, Any]) -> list[tuple[str, Any]]:
    """The report as ("h1"|"h2"|"p"|"table"|"figure", content) blocks."""
    out: list[tuple[str, Any]] = [
        ("title", spec.get("title") or spec.get("project", {}).get("name", "PropBench report"))
    ]
    if spec.get("authors"):
        out.append(("p", ", ".join(spec["authors"]) if isinstance(spec["authors"], list) else str(spec["authors"])))

    datasets = spec.get("datasets", [])
    out.append(("h1", "Data"))
    rows = []
    for d in datasets:
        t = d.get("temperature") or [0]
        p = [x for x in (d.get("pressure") or []) if x is not None]
        u = d.get("expanded_uncertainty") or []
        v = d.get("values") or []
        rel = [100 * a / b for a, b in zip(u, v, strict=False) if a is not None and b]
        prov = d.get("provenance") or {}
        rows.append(
            [
                d.get("name", ""),
                str(len(v)),
                f"{min(t):.1f}–{max(t):.1f}",
                f"{min(p) * 1e-6:.3g}–{max(p) * 1e-6:.3g}" if p else "saturation",
                f"{max(rel):.2g}" if rel else "–",
                prov.get("citation") or "",
            ]
        )
    out.append(("table", {"head": ["Dataset", "Points", "T / K", "p / MPa", "U / %", "Source"], "rows": rows}))

    cons = spec.get("consistency")
    if cons:
        out.append(("h1", "Data consistency"))
        for o in cons.get("overlaps", []):
            out.append(("p", f"{o['a']} and {o['b']} overlap at {o['t_range'][0]:.1f}–{o['t_range'][1]:.1f} K."))
        for w in cons.get("warnings", []):
            out.append(("p", f"Warning: {w}"))

    models = spec.get("models", [])
    if models:
        out.append(("h1", "Models"))
        for m in models:
            out.append(("h2", m.get("label", "model")))
            if m.get("reference"):
                out.append(("p", f"Model: {m['reference']}"))
            s = m.get("summary") or {}
            prow = [
                [name, _fmt(val, 6), _fmt((s.get("standard_errors") or {}).get(name), 3)]
                for name, val in (s.get("values") or {}).items()
            ]
            if prow:
                out.append(("table", {"head": ["Parameter", "Value", "Standard error"], "rows": prow}))
            dev = s.get("deviations") or {}
            if dev:
                out.append(
                    (
                        "p",
                        f"Fit: AARD {_fmt(dev.get('aard'))} %, bias {_fmt(dev.get('bias'))} %, "
                        f"max {_fmt(dev.get('max_ard'))} %.",
                    )
                )
            study = m.get("study") or {}
            cv_rows = [
                [
                    method,
                    str(len(cv.get("folds", []))),
                    _fmt((cv.get("pooled") or {}).get("aard")),
                    _fmt((cv.get("pooled") or {}).get("bias")),
                ]
                for method, cv in (study.get("cross_validation") or {}).items()
            ]
            if cv_rows:
                out.append(("table", {"head": ["Validation", "Folds", "AARD / %", "Bias / %"], "rows": cv_rows}))
            failed = [p["name"] for p in study.get("physics") or [] if not p.get("passed")]
            if study.get("physics") is not None:
                out.append(("p", "Physics checks: " + (", ".join(failed) + " failed." if failed else "all passed.")))

    comp = spec.get("comparison")
    if comp and comp.get("rows"):
        out.append(("h1", "Comparison of models"))
        out.append(
            (
                "table",
                {
                    "head": ["Model", "Dataset", "Points", "AARD / %", "Bias / %"],
                    "rows": [
                        [r["model"], r["dataset"], str(r.get("n", "")), _fmt(r.get("aard")), _fmt(r.get("bias"))]
                        for r in comp["rows"]
                    ],
                },
            )
        )

    out.append(("h1", "Reference models: check values"))
    out.append(
        (
            "table",
            {"head": ["Model", "State", "Published", "This installation", "Reproduced"], "rows": check_value_rows()},
        )
    )

    sel = spec.get("selection")
    if sel:
        out.append(("h1", "Model selection"))
        rule = sel.get("rule") or {}
        out.append(
            (
                "p",
                f"Selection rule locked before fitting (SHA-256 {str(sel.get('sha256', ''))[:16]}…): "
                f"metric {rule.get('metric')}, "
                f"validation {rule.get('validation')}. Chosen: {sel.get('chosen')} ({sel.get('reason', '')}).",
            )
        )

    for i, f in enumerate(spec.get("figures", []), start=1):
        out.append(("figure", {"n": i, "caption": f.get("caption", ""), "spec": f["spec"]}))

    refs = spec.get("references", [])
    if refs:
        out.append(("h1", "References"))
        for r in refs:
            out.append(("p", r.get("citation") or r.get("key", "")))

    out.append(("h1", "Reproducibility"))
    versions = {**_versions(), **(spec.get("versions") or {})}
    out.append(("p", "; ".join(f"{k} {v}" for k, v in versions.items()) + f". Seed {spec.get('seed', 0)}."))
    return out


# --- writers ---


def to_markdown(spec: Mapping[str, Any], figure_files: Sequence[str] = ()) -> str:
    lines: list[str] = []
    for kind, c in blocks(spec):
        if kind == "title":
            lines += [f"# {c}", ""]
        elif kind == "h1":
            lines += [f"## {c}", ""]
        elif kind == "h2":
            lines += [f"### {c}", ""]
        elif kind == "p":
            lines += [str(c), ""]
        elif kind == "table":
            lines.append("| " + " | ".join(c["head"]) + " |")
            lines.append("|" + "---|" * len(c["head"]))
            lines += ["| " + " | ".join(str(x).replace("|", "\\|") for x in row) + " |" for row in c["rows"]]
            lines.append("")
        elif kind == "figure":
            file = figure_files[c["n"] - 1] if len(figure_files) >= c["n"] else f"figure{c['n']}.svg"
            lines += [f"![Figure {c['n']}. {c['caption']}]({file})", "", f"*Figure {c['n']}.* {c['caption']}", ""]
    return "\n".join(lines)


def to_docx(spec: Mapping[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt

    from propbench import figures

    tpl = TEMPLATES.get(str(spec.get("template", "plain")), TEMPLATES["plain"])
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = tpl["docx_font"]
    style.font.size = Pt(tpl["size"])
    doc.core_properties.author = "PropBench"
    for kind, c in blocks(spec):
        if kind == "title":
            doc.add_heading(str(c), level=0)
        elif kind == "h1":
            doc.add_heading(str(c), level=1)
        elif kind == "h2":
            doc.add_heading(str(c), level=2)
        elif kind == "p":
            doc.add_paragraph(str(c))
        elif kind == "table":
            table = doc.add_table(rows=1, cols=len(c["head"]))
            table.style = "Light Grid Accent 1"
            for cell, text in zip(table.rows[0].cells, c["head"], strict=True):
                cell.text = str(text)
            for row in c["rows"]:
                cells = table.add_row().cells
                for cell, text in zip(cells, row, strict=False):
                    cell.text = str(text)
            doc.add_paragraph()
        elif kind == "figure":
            png = figures.render(c["spec"], "png", 300)
            doc.add_picture(io.BytesIO(png), width=Cm(14))
            doc.add_paragraph(f"Figure {c['n']}. {c['caption']}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _typ(text: Any) -> str:
    """Escape text for Typst markup."""
    s = str(text)
    for ch in "\\#$*_@<>[]`~=/":
        s = s.replace(ch, "\\" + ch)
    return s


def to_pdf(spec: Mapping[str, Any]) -> bytes:
    import typst

    from propbench import figures

    tpl = TEMPLATES.get(str(spec.get("template", "plain")), TEMPLATES["plain"])
    with tempfile.TemporaryDirectory() as tmp:
        parts = [
            f'#set page(paper: "a4", margin: 2.2cm)\n#set text(size: {tpl["size"]}pt)\n'
            "#set par(justify: true)\n#set table(stroke: 0.4pt)\n"
        ]
        for kind, c in blocks(spec):
            if kind == "title":
                parts.append(f'#align(center, text(size: 16pt, weight: "bold")[{_typ(c)}])\n')
            elif kind == "h1":
                parts.append(f"= {_typ(c)}\n")
            elif kind == "h2":
                parts.append(f"== {_typ(c)}\n")
            elif kind == "p":
                parts.append(f"{_typ(c)}\n")
            elif kind == "table":
                cells = ", ".join(f"[*{_typ(h)}*]" for h in c["head"])
                for row in c["rows"]:
                    cells += ", " + ", ".join(f"[{_typ(x)}]" for x in row)
                parts.append(f"#table(columns: {len(c['head'])}, {cells})\n")
            elif kind == "figure":
                name = f"figure{c['n']}.svg"
                (Path(tmp) / name).write_bytes(figures.render(c["spec"], "svg"))
                parts.append(f'#figure(image("{name}", width: 80%), caption: [{_typ(c["caption"])}])\n')
        source = Path(tmp) / "report.typ"
        source.write_text("\n".join(parts), encoding="utf-8")
        try:
            return typst.compile(str(source), root=tmp)
        except Exception as exc:  # typst raises its own error types
            raise ReportError(f"PDF could not be written: {exc}") from exc


def bundle(spec: Mapping[str, Any], project: Mapping[str, Any] | None = None) -> bytes:
    """Reproducibility bundle: the report (Markdown and PDF), figures, inputs, versions and seeds."""
    from propbench import figures

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        names = []
        for i, f in enumerate(spec.get("figures", []), start=1):
            name = f"figure{i}.svg"
            z.writestr(name, figures.render(f["spec"], "svg"))
            z.writestr(f"figure{i}.pdf", figures.render(f["spec"], "pdf"))
            names.append(name)
        z.writestr("report.md", to_markdown(spec, names))
        z.writestr("report.pdf", to_pdf(spec))
        z.writestr("inputs/report_spec.json", json.dumps(spec, indent=1, default=str))
        if project is not None:
            z.writestr("inputs/project.json", json.dumps(project, indent=1, default=str))
        z.writestr("versions.json", json.dumps({**_versions(), "seed": spec.get("seed", 0)}, indent=1))
    return buf.getvalue()


def render(spec: Mapping[str, Any], fmt: str, project: Mapping[str, Any] | None = None) -> dict[str, Any]:
    if fmt not in FORMATS:
        raise ReportError(f"format must be one of {sorted(FORMATS)}")
    if fmt == "md":
        data = to_markdown(spec).encode("utf-8")
    elif fmt == "docx":
        data = to_docx(spec)
    elif fmt == "pdf":
        data = to_pdf(spec)
    else:
        data = bundle(spec, project)
    return {
        "format": fmt,
        "mime": FORMATS[fmt],
        "content_base64": base64.b64encode(data).decode("ascii"),
        "bytes": len(data),
    }
