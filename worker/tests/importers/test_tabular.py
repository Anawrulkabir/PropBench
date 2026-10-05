import numpy as np
import pytest
from openpyxl import Workbook

from propbench.core import Phase, Provenance, Quantity
from propbench.importers import (
    ColumnSpec,
    ImportFileError,
    ImportMapping,
    MappingError,
    Role,
    UncertaintyKind,
    import_file,
)


def viscosity_mapping(**overrides):
    kwargs = {
        "quantity": Quantity.VISCOSITY,
        "columns": (
            ColumnSpec("T", Role.TEMPERATURE, "K"),
            ColumnSpec("p", Role.PRESSURE, "MPa"),
            ColumnSpec("eta", Role.VALUE, "μPa·s"),
            ColumnSpec("U", Role.UNCERTAINTY, uncertainty_kind=UncertaintyKind.RELATIVE_PERCENT),
        ),
    }
    kwargs.update(overrides)
    return ImportMapping(**kwargs)


def test_csv_with_units_row_relative_uncertainty_and_provenance(tmp_path):
    path = tmp_path / "cf3i_viscosity.csv"
    path.write_text(
        "T,p,eta,U\nK,MPa,μPa·s,%\n300.0,1.0,201.5,2.0\n\n310.0,5.0,190.0,2.5\n",
        encoding="utf-8",
    )
    mapping = viscosity_mapping(first_data_row=3)
    ds = import_file(path, mapping, fluid="CF3I", provenance=Provenance(doi="10.1000/x", method="vibrating wire"))
    assert ds.name == "cf3i_viscosity"
    assert ds.fluid == "R13I1"
    np.testing.assert_allclose(ds.temperature, [300.0, 310.0])
    np.testing.assert_allclose(ds.pressure, [1e6, 5e6])
    np.testing.assert_allclose(ds.values, [201.5e-6, 190.0e-6])
    np.testing.assert_allclose(ds.expanded_uncertainty, [201.5e-6 * 0.02, 190.0e-6 * 0.025])
    assert ds.coverage_factor == 2.0
    assert ds.provenance.method == "vibrating wire"
    np.testing.assert_array_equal(ds.point_ids, [0, 1])  # the blank line is skipped


def test_semicolon_decimal_comma_celsius_and_absolute_uncertainty(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("T/°C;p/bar;rho;u\n26,85;10;1201,53;0,5\n6,85;20;1274,52;0,5\n", encoding="utf-8-sig")
    mapping = ImportMapping(
        quantity=Quantity.MASS_DENSITY,
        columns=(
            ColumnSpec("T/°C", Role.TEMPERATURE, "°C"),
            ColumnSpec("p/bar", Role.PRESSURE, "bar"),
            ColumnSpec("rho", Role.VALUE, "kg/m³"),
            ColumnSpec("u", Role.UNCERTAINTY, "kg/m³", UncertaintyKind.ABSOLUTE),
        ),
        decimal_comma=True,
        coverage_factor=1.0,
    )
    ds = import_file(path, mapping, fluid="R134a")
    np.testing.assert_allclose(ds.temperature, [300.0, 280.0])
    np.testing.assert_allclose(ds.pressure, [1e6, 2e6])
    np.testing.assert_allclose(ds.values, [1201.53, 1274.52])
    np.testing.assert_allclose(ds.expanded_uncertainty, [0.5, 0.5])
    assert ds.coverage_factor == 1.0


def test_xlsx_named_sheet_with_phase_column(tmp_path):
    path = tmp_path / "Viscosity_CF3I.xlsx"
    workbook = Workbook()
    workbook.active.title = "notes"
    sheet = workbook.create_sheet("data")
    sheet.append(["T", "p", "eta", "U", "phase"])
    sheet.append([300, 1.0, 201.5, 2.0, "L"])
    sheet.append([350.0, 0.5, 15.1, 3.0, "vapour"])
    workbook.save(path)
    mapping = viscosity_mapping(
        sheet="data",
        columns=(*viscosity_mapping().columns, ColumnSpec("phase", Role.PHASE)),
    )
    ds = import_file(path, mapping, fluid="2314-97-8")
    assert ds.fluid == "R13I1"
    assert ds.phase == (Phase.LIQUID, Phase.VAPOR)
    np.testing.assert_allclose(ds.values, [201.5e-6, 15.1e-6])


def test_xlsx_unknown_sheet(tmp_path):
    path = tmp_path / "book.xlsx"
    Workbook().save(path)
    with pytest.raises(ImportFileError, match="no sheet 'data'"):
        import_file(path, viscosity_mapping(sheet="data"), fluid="R134a")


@pytest.mark.parametrize(
    ("content", "message"),
    [
        ("T,p,eta,U\n300,1,abc,2\n", "row 2, column 'eta': 'abc' is not a number"),
        ("T,p,eta,U\n300,,201,2\n", "row 2, column 'p': empty cell"),
        ("T,pressure,eta,U\n300,1,201,2\n", "no column 'p'"),
        ("T,p,eta,U\n", "no data rows"),
        ("T,p,eta,U\n-300,1,201,2\n", "temperatures must be positive"),
    ],
)
def test_errors_name_the_row_and_column(tmp_path, content, message):
    path = tmp_path / "bad.csv"
    path.write_text(content, encoding="utf-8")
    with pytest.raises(ImportFileError, match=message):
        import_file(path, viscosity_mapping(), fluid="R134a")


def test_unknown_phase_word(tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text("T,p,eta,U,phase\n300,1,201,2,plasma\n", encoding="utf-8")
    mapping = viscosity_mapping(columns=(*viscosity_mapping().columns, ColumnSpec("phase", Role.PHASE)))
    with pytest.raises(ImportFileError, match="unknown phase 'plasma'"):
        import_file(path, mapping, fluid="R134a")


def test_unsupported_file_type(tmp_path):
    path = tmp_path / "data.ods"
    path.write_text("")
    with pytest.raises(ImportFileError, match="unsupported file type"):
        import_file(path, viscosity_mapping(), fluid="R134a")


def test_mapping_template_round_trip(tmp_path):
    mapping = viscosity_mapping(first_data_row=3, sheet="data", notes="Tuhin 2024 layout")
    path = tmp_path / "template.json"
    mapping.save(path)
    assert ImportMapping.load(path) == mapping
    assert '"kind": "propbench.import-mapping"' in path.read_text(encoding="utf-8")


@pytest.mark.parametrize(
    ("change", "message"),
    [
        ({"columns": (ColumnSpec("T", Role.TEMPERATURE, "K"), ColumnSpec("p", Role.PRESSURE, "MPa"))}, "'value'"),
        (
            {"columns": (ColumnSpec("T", Role.TEMPERATURE, "K"), ColumnSpec("eta", Role.VALUE, "Pa*s"))},
            "pressure or molar-density",
        ),
        (
            {
                "columns": (
                    ColumnSpec("T", Role.TEMPERATURE, "K"),
                    ColumnSpec("p", Role.PRESSURE, "MPa"),
                    ColumnSpec("eta", Role.VALUE, "MPa"),
                )
            },
            "not a unit of viscosity",
        ),
        (
            {
                "columns": (
                    ColumnSpec("T", Role.TEMPERATURE),
                    ColumnSpec("p", Role.PRESSURE, "MPa"),
                    ColumnSpec("eta", Role.VALUE, "Pa*s"),
                )
            },
            "needs a unit",
        ),
        (
            {
                "columns": (
                    ColumnSpec("T", Role.TEMPERATURE, "K"),
                    ColumnSpec("p", Role.PRESSURE, "MPa"),
                    ColumnSpec("eta", Role.VALUE, "Pa*s"),
                    ColumnSpec("U", Role.UNCERTAINTY),
                )
            },
            "needs an uncertainty_kind",
        ),
        ({"header_row": 0}, "1-based"),
        ({"first_data_row": 1}, "after header_row"),
    ],
)
def test_mapping_validation(change, message):
    with pytest.raises((MappingError, ValueError), match=message):
        viscosity_mapping(**change)


def test_template_of_another_kind_is_rejected():
    with pytest.raises(MappingError, match="not a version-1"):
        ImportMapping.from_dict({"kind": "something-else", "version": 1})
